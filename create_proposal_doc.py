import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import comtypes.client

def create_document():
    doc = docx.Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Header with Official Word Watermark
        header = section.header
        # Replace the default paragraph in header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        watermark_xml = (
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:v="urn:schemas-microsoft-com:vml"'
            ' xmlns:o="urn:schemas-microsoft-com:office:office">'
            '<w:pPr><w:pStyle w:val="Header"/></w:pPr>'
            '<w:r>'
            '<w:rPr><w:noProof/></w:rPr>'
            '<w:pict>'
            '<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">'
            '<v:formulas>'
            '<v:f eqn="val #0"/><v:f eqn="prod #0 4 3"/><v:f eqn="prod #0 1 3"/>'
            '<v:f eqn="sum @1 0 21600"/><v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 21600 0 @3"/>'
            '<v:f eqn="sum @4 21600 0"/><v:f eqn="prod @2 1 2"/><v:f eqn="prod @3 1 2"/>'
            '<v:f eqn="val 10800"/><v:f eqn="prod 10800 1 2"/>'
            '</v:formulas>'
            '<v:path textpathok="t" o:connecttype="rect"/>'
            '<v:textpath on="t" fitshape="t"/>'
            '<o:lock v:ext="edit" text="t" shapetype="t"/>'
            '</v:shapetype>'
            '<v:shape id="PowerPlusWaterMarkObject357476642" o:spid="_x0000_s1025" type="#_x0000_t136"'
            ' style="position:absolute;margin-left:0;margin-top:0;width:480pt;height:125pt;rotation:315;z-index:-251656704;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"'
            ' fillcolor="#cbd5e1" stroked="f">'
            '<v:fill opacity="0.22"/>'
            '<v:textpath style="font-family:\'Arial\',sans-serif;font-size:1pt;font-weight:bold" string="NexoWebs"/>'
            '</v:shape>'
            '</w:pict>'
            '</w:r>'
            '</w:p>'
        )
        wp = parse_xml(watermark_xml)
        # Clear default p and append our watermark p
        header._element.remove(hp._p)
        header._element.append(wp)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("NexoWebs · Soluciones Web y Cartas Digitales para el Sector Gastronómico")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(140, 145, 155)

    # Styles
    NAVY = RGBColor(15, 35, 65)
    BLUE = RGBColor(30, 95, 170)
    DARK = RGBColor(35, 40, 50)
    GRAY = RGBColor(95, 105, 115)
    
    def add_title(text, subtitle=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(21)
        run.font.color.rgb = NAVY
        
        if subtitle:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(8)
            run2 = p2.add_run(subtitle)
            run2.bold = True
            run2.font.name = "Calibri"
            run2.font.size = Pt(12)
            run2.font.color.rgb = BLUE

    def add_meta_banner(client_text, presenter_text):
        tbl = doc.add_table(rows=1, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_l = tbl.cell(0, 0)
        cell_r = tbl.cell(0, 1)
        
        shd_xml = '<w:shd {} w:val="clear" w:color="auto" w:fill="F1F5F9"/>'.format(nsdecls('w'))
        cell_l._tc.get_or_add_tcPr().append(parse_xml(shd_xml))
        cell_r._tc.get_or_add_tcPr().append(parse_xml(shd_xml))
        
        p_l = cell_l.paragraphs[0]
        p_l.paragraph_format.space_before = Pt(5)
        p_l.paragraph_format.space_after = Pt(5)
        r_l_bold = p_l.add_run("Dirigido a: ")
        r_l_bold.bold = True
        r_l_bold.font.size = Pt(9.5)
        r_l_bold.font.color.rgb = DARK
        r_l_val = p_l.add_run(client_text)
        r_l_val.font.size = Pt(9.5)
        r_l_val.font.color.rgb = DARK
        
        p_r = cell_r.paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_r.paragraph_format.space_before = Pt(5)
        p_r.paragraph_format.space_after = Pt(5)
        r_r_bold = p_r.add_run("Presentado por: ")
        r_r_bold.bold = True
        r_r_bold.font.size = Pt(9.5)
        r_r_bold.font.color.rgb = DARK
        r_r_val = p_r.add_run(presenter_text)
        r_r_val.bold = True
        r_r_val.font.size = Pt(9.5)
        r_r_val.font.color.rgb = BLUE
        
        borders_xml = (
            '<w:tcBorders {} >'
            '<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
            '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>'
            '<w:left w:val="none"/>'
            '<w:right w:val="none"/>'
            '</w:tcBorders>'
        ).format(nsdecls('w'))
        cell_l._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
        cell_r._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.color.rgb = NAVY
        
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(
            '<w:pBdr {}><w:bottom w:val="single" w:sz="10" w:space="2" w:color="2563EB"/></w:pBdr>'.format(nsdecls('w'))
        )
        pPr.append(pBdr)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(11.5)
        run.font.color.rgb = BLUE

    def add_body(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(10)
            r_bold.font.color.rgb = DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = DARK

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.bold = True
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(10)
            r_bold.font.color.rgb = DARK
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = DARK

    def add_tip_box(text, prefix="💡 Recomendación para tu negocio: "):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        
        shd_xml = '<w:shd {} w:val="clear" w:color="auto" w:fill="FEF9C3"/>'.format(nsdecls('w'))
        cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))
        
        border_xml = (
            '<w:tcBorders {} >'
            '<w:top w:val="none"/>'
            '<w:bottom w:val="none"/>'
            '<w:left w:val="single" w:sz="24" w:space="0" w:color="CA8A04"/>'
            '<w:right w:val="none"/>'
            '</w:tcBorders>'
        ).format(nsdecls('w'))
        cell._tc.get_or_add_tcPr().append(parse_xml(border_xml))
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.line_spacing = 1.15
        
        r_pre = p.add_run(prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor(133, 77, 14)
        
        r_txt = p.add_run(text)
        r_txt.font.name = "Calibri"
        r_txt.font.size = Pt(9.5)
        r_txt.font.color.rgb = RGBColor(60, 40, 10)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_image_block(img_filename, caption_text, width_inches=5.8):
        img_path = os.path.join("extracted_images", img_filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.add_run().add_picture(img_path, width=Inches(width_inches))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(0)
            p_cap.paragraph_format.space_after = Pt(8)
            rcap = p_cap.add_run(caption_text)
            rcap.italic = True
            rcap.font.name = "Calibri"
            rcap.font.size = Pt(8.5)
            rcap.font.color.rgb = GRAY

    # --- DOCUMENT GENERATION ---
    add_title("PROPUESTA DE SERVICIO WEB Y CARTA DIGITAL", "Solución Digital Profesional para Empresas y Locales Gastronómicos")
    add_meta_banner("Empresas y Locales del Rubro Gastronómico", "Equipo NexoWebs")

    add_body(
        "Hemos preparado este documento de forma clara y directa para mostrarte cómo desarrollamos la página web oficial "
        "y la carta digital interactiva para tu negocio gastronómico (ej. www.tunegocio.cl). Nuestro objetivo es que tus clientes "
        "puedan explorar toda tu oferta de platos, especialidades, combos y promociones desde su teléfono celular o computador "
        "en segundos, y enviarte sus pedidos de manera directa a tu WhatsApp comercial, sin intermediarios ni cobro "
        "de comisiones por cada venta."
    )

    # 1. QUE INCLUYE
    add_h1("1. ¿Qué incluye exactamente el servicio que te ofrecemos?")
    
    add_bullet(
        "Diseñamos tu sitio web respetando los colores, logotipo, estilo visual y la identidad propia de tu negocio de comida. "
        "La interfaz está 100% optimizada para teléfonos móviles (desde donde ingresa la gran mayoría de tus comensales a través "
        "de redes sociales como Instagram o WhatsApp) y computadores, presentando tus platos y bebidas de forma apetitosa, rápida y ordenada.",
        bold_prefix="Creación y Diseño Atractivo a la Medida de tu Marca: "
    )
    add_bullet(
        "Registramos y dejamos activo tu propio enlace web oficial (.cl) incluido dentro del valor de tu plan inicial ($50.000 o $60.000), "
        "para que lo incorpores en tu perfil de Instagram, estados de WhatsApp, flyers, empaques, bolsas de delivery o códigos QR en las mesas de tu local.",
        bold_prefix="Tu Propio Nombre en Internet (Dominio .CL Incluido por 1 Año): "
    )
    add_bullet(
        "Sabemos que en un negocio de comida los precios, ingredientes o promociones cambian constantemente. "
        "Por eso, contarás con un apartado privado (con tu usuario y contraseña exclusivos) donde podrás añadir productos nuevos, "
        "modificar precios, subir fotografías o pausar platos agotados cuando quieras, directamente desde tu propio celular o computador.",
        bold_prefix="Apartado Exclusivo para que Administres tu Menú: "
    )
    add_bullet(
        "Tus clientes navegan por las categorías de tu carta, seleccionan lo que desean y, con un solo toque, el sistema genera "
        "un mensaje claro con el pedido desglosado directo a tu número de WhatsApp comercial, además de integrar la ubicación exacta de tu local "
        "en Google Maps y tus horarios de apertura y cierre.",
        bold_prefix="Pedidos Directos a tu WhatsApp Comercial: "
    )
    add_bullet(
        "Aunque la plataforma está diseñada para ser sumamente intuitiva y fácil de usar por tu cuenta, nunca estarás solo. "
        "Siempre que necesites ayuda para subir un producto, actualizar fotos o realizar cualquier ajuste, "
        "nosotros estaremos totalmente disponibles para apoyarte.",
        bold_prefix="Acompañamiento y Soporte Técnico Continuo: "
    )

    # 2. COMO ADMINISTRAR
    add_h1("2. ¿Cómo podrás añadir o eliminar productos de tu Carta?")
    add_body(
        "No necesitarás saber nada de programación ni diseño técnico. Te entregaremos un Panel de Administración "
        "muy simple y conectado a la base de datos de tu página web:"
    )
    
    add_body("Desde tu celular o computador ingresas a tu apartado de dueño con tu usuario y contraseña.", bold_prefix="1. Entras con tu clave secreta: ")
    add_body("Escribes el nombre del plato o combo, eliges la categoría (ej. Platos Fuertes, Promociones, Entradas, Postres, Bebidas), pones el precio, una breve descripción y subes la fotografía.", bold_prefix="2. Añadir un plato o promoción: ")
    add_body("Si subió un insumo o se acabó un plato por la noche, con un solo toque puedes cambiarle el precio, pausarlo temporalmente o eliminarlo de la carta.", bold_prefix="3. Modificar o pausar un producto: ")
    add_body("Apenas guardas el cambio, todos tus clientes ya ven el menú actualizado al instante en tu sitio web.", bold_prefix="4. Cambio instantáneo en tiempo real: ")
    add_body("Si en algún turno estás ocupado en la cocina o atendiendo el local y necesitas que nosotros cambiemos un precio o agreguemos una promoción por ti, nos escribes y nosotros lo hacemos.", bold_prefix="5. ¿Poco tiempo en el local? ¡Nosotros te ayudamos!: ")

    # 3. PLANES Y PRECIOS
    add_h1("3. Nuestros Planes y Precios (Con Dominio .CL Incluido)")
    add_body("Hemos diseñado dos opciones claras para que elijas la que mejor se adapte a lo que buscas para tu negocio gastronómico:")

    # Table of Plans
    table = doc.add_table(rows=8, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(3.3), Inches(1.8), Inches(2.1)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    headers = [
        "¿Qué incluye cada opción?",
        "PLAN 1: CARTA DIGITAL\n$50.000 CLP",
        "PLAN 2: WEB COMPLETA + CARTA\n$60.000 CLP (Recomendado)"
    ]
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        
        bg_color = "1E3A8A" if idx < 2 else "1D4ED8"
        shd_xml = '<w:shd {} w:val="clear" w:color="auto" w:fill="{}"/>'.format(nsdecls('w'), bg_color)
        cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))

    rows_data = [
        ("Dominio propio en Internet (.CL - 1er Año)", "✅ Incluido en los $50.000", "✅ Incluido en los $60.000"),
        ("Carta Digital Interactiva (fotos, categorías y precios)", "✅ Sí, completa", "✅ Sí, completa"),
        ("Apartado del Dueño (Añadir / Editar / Pausar platos)", "✅ Sí, autoadministrable", "✅ Sí, autoadministrable"),
        ("Página Principal de Bienvenida (Landing Page + Sobre Nosotros + Historia + Mapa)", "❌ No (entra directo a la Carta)", "✅ Sí, Sitio Web Institucional Completo"),
        ("Botones directos a WhatsApp y Redes Sociales", "✅ Sí", "✅ Sí"),
        ("Integración y configuración completa al Hosting", "✅ Sí, lista para operar", "✅ Sí, lista para operar"),
        ("Soporte y acompañamiento continuo", "✅ Siempre disponibles", "✅ Siempre disponibles"),
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(9)
            if col_idx > 0 and "✅" in text:
                run.font.color.rgb = RGBColor(21, 128, 61)
            elif col_idx > 0 and "❌" in text:
                run.font.color.rgb = RGBColor(185, 28, 28)
            else:
                run.font.color.rgb = DARK
                
            if row_idx % 2 == 1:
                shd_xml = '<w:shd {} w:val="clear" w:color="auto" w:fill="F8FAFC"/>'.format(nsdecls('w'))
                cell._tc.get_or_add_tcPr().append(parse_xml(shd_xml))

    tblPr = table._tbl.tblPr
    borders_xml = (
        '<w:tblBorders {} >'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        '</w:tblBorders>'
    ).format(nsdecls('w'))
    tblPr.append(parse_xml(borders_xml))

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_tip_box(
        "Por solo $10.000 de diferencia, el Plan de $60.000 te entrega no solo la carta interactiva, sino también la Página Principal "
        "completa con la presentación de tu negocio, tu historia ('Sobre Nosotros'), vitrina de especialidades destacadas, "
        "mapa de ubicación interactivo y horarios de atención en vivo."
    )

    # 4. TRANSPARENCIA TOTAL
    add_h1("4. Transparencia Total: Infraestructura de Dominio .CL y Hosting")
    add_body(
        "Creemos en una relación de total transparencia y confianza con cada local. A continuación detallamos los valores oficiales "
        "y especificaciones de los servicios técnicos donde se aloja tu plataforma (SolucionHost y registro de dominios):"
    )

    add_h2("A. Costo del Dominio Web (.CL) según cantidad de años")
    add_body(
        "El primer año de tu dominio oficial ($8.900 + IVA) ya viene 100% incluido dentro del valor de tu plan ($50.000 o $60.000). "
        "Para que tengas total claridad hacia el futuro (o por si deseas dejarlo asegurado por más años desde el inicio), "
        "estos son los valores oficiales de registro según la cantidad de años:"
    )

    add_image_block(
        "image1.png",
        "Imagen 1: Tarifas oficiales de registro de dominio .cl según los años de registro (1 Año: $8.900 | 2 Años: $17.800 | 3 Años: $26.700 | 4 Años: $35.600 | 5 Años: $44.500 neto).",
        width_inches=5.2
    )

    add_h2("B. Servicio de Hosting y Base de Datos (Plan Estándar 1)")
    add_body(
        "Para que tu página web esté encendida las 24 horas del día y guarde todos los productos y fotos que subas a tu menú, "
        "nosotros dejamos todo integrado en el servicio de hosting profesional SolucionHost (Plan Estándar 1):"
    )

    add_bullet(
        "Tu dirección de internet oficial conectada directamente al servidor de alta disponibilidad.",
        bold_prefix="Dominio Hospedado: "
    )
    add_bullet(
        "Capacidad de sobra para guardar todas las fotografías de tus platos, promociones y diseño de la página con carga rápida.",
        bold_prefix="20 GB de Espacio Web: "
    )
    add_bullet(
        "Aquí es donde se guarda tu menú digital para que cada vez que agregues o elimines un plato desde tu apartado de dueño, el cambio quede guardado al instante.",
        bold_prefix="2 Bases de Datos Incluidas: "
    )
    add_bullet(
        "Incluye la posibilidad de contar con casillas de correos con el nombre de tu propio local (ej. contacto@tunegocio.cl).",
        bold_prefix="Cuentas de Correo Ilimitadas: "
    )
    add_bullet(
        "Puedes elegir ciclo Anual por $11.900 + IVA ($14.161 total al año), Bi-Anual por $18.900 + IVA, o Trienal por $20.900 + IVA, como se aprecia en las capturas oficiales:",
        bold_prefix="Ciclos de Facturación del Hosting en SolucionHost: "
    )

    add_image_block(
        "image2.png",
        "Imagen 2: Selección del dominio y ciclos de facturación del Hosting Plan Estándar 1.",
        width_inches=5.5
    )

    add_image_block(
        "image3.png",
        "Imagen 3: Detalle de capacidad técnica incluida en el Plan Estándar 1 (20 GB Web, 2 Bases de Datos y cuentas de correo).",
        width_inches=2.8
    )

    add_h2("C. Resumen del Total Anual (Hosting + Dominio + IVA)")
    add_body(
        "En la siguiente imagen puedes ver el desglose exacto y transparente de ambos servicios juntos en SolucionHost: "
        "Hosting Plan Estándar 1 ($11.900) + Registro de Dominio por 1 Año ($8.900) + IVA ($3.952) = $24.752 CLP Anual en total "
        "(recuerda que este costo inicial ya va incluido dentro de tu plan de $50.000 o $60.000):"
    )

    add_image_block(
        "image4.png",
        "Imagen 4: Total oficial en SolucionHost sumando Hosting Estándar 1 ($11.900) + Dominio .cl por 1 año ($8.900) + IVA = $24.752 CLP Anual.",
        width_inches=5.5
    )

    # 5. COMO EMPEZAMOS
    add_h1("5. ¿Cómo empezamos y cómo te acompañamos?")
    add_body("Diseñamos un proceso ágil en 4 pasos para poner en marcha tu solución web sin fricción:")

    add_body("Nos confirmas si deseas comenzar con el Plan 1 ($50.000 solo Carta Digital + Dominio) o el Plan 2 ($60.000 Página Principal Completa + Sobre Nosotros + Carta Digital + Dominio).", bold_prefix="Paso 1 - Eliges tu Plan: ")
    add_body("Dejamos registrado el dominio .cl, conectamos el hosting, armamos la base de datos y subimos todos tus productos actuales para que no tengas que partir desde cero.", bold_prefix="Paso 2 - Nosotros configuramos todo por ti: ")
    add_body("Te mostramos la página funcionando para que revises cada detalle y nos digas si quieres ajustar algún texto, color, precio o fotografía.", bold_prefix="Paso 3 - Revisión junto a ti: ")
    add_body("Te enseñamos en pocos minutos cómo agregar o borrar productos de tu menú. Y recuerda: cada vez que necesites ayuda o prefieras que nosotros hagamos un cambio por ti, estaremos disponibles para apoyarte.", bold_prefix="Paso 4 - Entrega de tu Panel y Soporte Continuo: ")

    # Closing
    p_close = doc.add_paragraph()
    p_close.paragraph_format.space_before = Pt(14)
    p_close.paragraph_format.space_after = Pt(2)
    p_close.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_close = p_close.add_run("¡Estamos listos para llevar la propuesta gastronómica de tu negocio a su propia página web oficial!")
    r_close.bold = True
    r_close.font.name = "Calibri"
    r_close.font.size = Pt(11)
    r_close.font.color.rgb = NAVY

    p_team = doc.add_paragraph()
    p_team.paragraph_format.space_before = Pt(0)
    p_team.paragraph_format.space_after = Pt(10)
    p_team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_team = p_team.add_run("Equipo NexoWebs — Siempre disponibles para ayudarte.")
    r_team.font.name = "Calibri"
    r_team.font.size = Pt(10)
    r_team.font.color.rgb = BLUE

    # Save DOCX
    docx_filename = "Propuesta_Servicio_Web_Gastronomico_NexoWebs.docx"
    doc.save(docx_filename)
    print(f"Document saved successfully as {docx_filename}")

    # Also export PDF copy using Word COM for immediate presentation
    try:
        word = comtypes.client.CreateObject('Word.Application')
        word.Visible = False
        full_docx = os.path.abspath(docx_filename)
        full_pdf = os.path.abspath("Propuesta_Servicio_Web_Gastronomico_NexoWebs.pdf")
        wdoc = word.Documents.Open(full_docx)
        wdoc.SaveAs(full_pdf, FileFormat=17) # 17 = wdFormatPDF
        wdoc.Close()
        word.Quit()
        print(f"PDF copy generated successfully at {full_pdf}")
    except Exception as e:
        print(f"Note on PDF export: {e}")

if __name__ == "__main__":
    create_document()
